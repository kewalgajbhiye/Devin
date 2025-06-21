"""
Export utilities for PDF, Excel, and CSV generation
"""

import os
import csv
from datetime import datetime
from tkinter import filedialog, messagebox

try:
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

class ExportUtils:
    @staticmethod
    def export_to_csv(data, headers, filename_prefix="export"):
        """Export data to CSV file"""
        try:
            filename = filedialog.asksaveasfilename(
                defaultextension=".csv",
                filetypes=[("CSV files", "*.csv")],
                initialfile=f"{filename_prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
            )
            
            if filename:
                with open(filename, 'w', newline='', encoding='utf-8') as csvfile:
                    writer = csv.writer(csvfile)
                    writer.writerow(headers)
                    writer.writerows(data)
                
                messagebox.showinfo("Success", f"Data exported successfully to {os.path.basename(filename)}")
                return True
            return False
            
        except Exception as e:
            messagebox.showerror("Error", f"Failed to export CSV: {e}")
            return False
    
    @staticmethod
    def export_to_excel(data, headers, filename_prefix="export", sheet_name="Data"):
        """Export data to Excel file"""
        if not OPENPYXL_AVAILABLE:
            messagebox.showerror("Error", "Excel export requires openpyxl. Install with: pip install openpyxl")
            return False
        
        try:
            filename = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
                initialfile=f"{filename_prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            )
            
            if filename:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = sheet_name
                
                for col, header in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=col, value=header)
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
                
                for row, row_data in enumerate(data, 2):
                    for col, value in enumerate(row_data, 1):
                        ws.cell(row=row, column=col, value=value)
                
                for column in ws.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    ws.column_dimensions[column_letter].width = adjusted_width
                
                wb.save(filename)
                messagebox.showinfo("Success", f"Data exported successfully to {os.path.basename(filename)}")
                return True
            return False
            
        except Exception as e:
            messagebox.showerror("Error", f"Failed to export Excel: {e}")
            return False
    
    @staticmethod
    def export_to_pdf(data, headers, title, filename_prefix="export"):
        """Export data to PDF file"""
        if not REPORTLAB_AVAILABLE:
            messagebox.showerror("Error", "PDF export requires reportlab. Install with: pip install reportlab")
            return False
        
        try:
            filename = filedialog.asksaveasfilename(
                defaultextension=".pdf",
                filetypes=[("PDF files", "*.pdf")],
                initialfile=f"{filename_prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
            )
            
            if filename:
                doc = SimpleDocTemplate(filename, pagesize=A4)
                elements = []
                styles = getSampleStyleSheet()
                
                title_style = ParagraphStyle(
                    'CustomTitle',
                    parent=styles['Heading1'],
                    fontSize=16,
                    spaceAfter=30,
                    alignment=1
                )
                elements.append(Paragraph(title, title_style))
                elements.append(Spacer(1, 12))
                
                table_data = [headers] + data
                table = Table(table_data)
                
                table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 12),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black)
                ]))
                
                elements.append(table)
                doc.build(elements)
                
                messagebox.showinfo("Success", f"Data exported successfully to {os.path.basename(filename)}")
                return True
            return False
            
        except Exception as e:
            messagebox.showerror("Error", f"Failed to export PDF: {e}")
            return False
